"""
ChronoMate Desktop - Export Engine
Exports session datasets to PDF, CSV, and Excel (XLSX).
"""

import csv
from datetime import datetime
from pathlib import Path
from typing import List, Dict, Any, Optional

from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.graphics.shapes import Drawing, Rect, Line, Circle, String

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

from core.models import ChronoSession, Shot


class Exporter:
    @staticmethod
    def export_csv(file_path: Path, shots: List[Shot], gun_name: str = "", player_name: str = "") -> bool:
        try:
            with open(file_path, "w", newline="", encoding="utf-8") as f:
                writer = csv.writer(f)
                writer.writerow(["#", "Timestamp", "Weight (g)", "Velocity (m/s)", "Energy (J)", "Comp Factor"])
                for idx, shot in enumerate(shots, 1):
                    dt_str = datetime.fromtimestamp(shot.timestamp).strftime("%Y-%m-%d %H:%M:%S")
                    writer.writerow([
                        idx,
                        dt_str,
                        f"{shot.weight_grams:.2f}",
                        f"{shot.velocity:.2f}",
                        f"{shot.energy_joules:.2f}",
                        f"{shot.compensation_factor:.3f}",
                    ])
            return True
        except Exception as e:
            print(f"CSV export error: {e}")
            return False

    @staticmethod
    def export_xlsx(file_path: Path, shots: List[Shot], gun_name: str = "", player_name: str = "", stats: Optional[Dict[str, Any]] = None) -> bool:
        if not OPENPYXL_AVAILABLE:
            return False

        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "ChronoMate Session"

            # Colors & Fonts
            header_fill = PatternFill(start_color="1B2838", end_color="1B2838", fill_type="solid")
            header_font = Font(name="Segoe UI", size=11, bold=True, color="88FF11")
            meta_font = Font(name="Segoe UI", size=10, bold=True, color="333333")
            data_font = Font(name="Segoe UI", size=10)

            # Metadata
            ws["A1"] = "ChronoMate Shooting Session Report"
            ws["A1"].font = Font(name="Segoe UI", size=16, bold=True, color="1B5E20")
            ws["A2"] = f"Date: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
            ws["A2"].font = meta_font
            if gun_name:
                ws["A3"] = f"Gun: {gun_name}"
                ws["A3"].font = meta_font
            if player_name:
                ws["A4"] = f"Player: {player_name}"
                ws["A4"].font = meta_font

            # Stats Summary Row
            start_row = 6
            if stats:
                ws.cell(row=start_row, column=1, value="AVERAGE").font = header_font
                ws.cell(row=start_row, column=1).fill = header_fill
                ws.cell(row=start_row + 1, column=1, value=f"{stats.get('avg', 0.0):.1f} m/s")

                ws.cell(row=start_row, column=2, value="MAX").font = header_font
                ws.cell(row=start_row, column=2).fill = header_fill
                ws.cell(row=start_row + 1, column=2, value=f"{stats.get('max', 0.0):.1f} m/s")

                ws.cell(row=start_row, column=3, value="MIN").font = header_font
                ws.cell(row=start_row, column=3).fill = header_fill
                ws.cell(row=start_row + 1, column=3, value=f"{stats.get('min', 0.0):.1f} m/s")

                ws.cell(row=start_row, column=4, value="EXTREME SPREAD").font = header_font
                ws.cell(row=start_row, column=4).fill = header_fill
                ws.cell(row=start_row + 1, column=4, value=f"{stats.get('es', 0.0):.1f} m/s")

                ws.cell(row=start_row, column=5, value="STD DEV").font = header_font
                ws.cell(row=start_row, column=5).fill = header_fill
                ws.cell(row=start_row + 1, column=5, value=f"{stats.get('sd', 0.0):.2f} m/s")

                ws.cell(row=start_row, column=6, value="ROF").font = header_font
                ws.cell(row=start_row, column=6).fill = header_fill
                ws.cell(row=start_row + 1, column=6, value=f"{stats.get('rof', '0.0')} r/m")

                start_row += 3

            # Table Header
            headers = ["#", "Timestamp", "Weight (g)", "Velocity (m/s)", "Energy (J)", "Comp Factor"]
            for col_idx, h in enumerate(headers, 1):
                cell = ws.cell(row=start_row, column=col_idx, value=h)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center")

            # Data Rows
            for idx, shot in enumerate(shots, 1):
                row_idx = start_row + idx
                dt_str = datetime.fromtimestamp(shot.timestamp).strftime("%H:%M:%S")
                ws.cell(row=row_idx, column=1, value=idx).font = data_font
                ws.cell(row=row_idx, column=2, value=dt_str).font = data_font
                ws.cell(row=row_idx, column=3, value=shot.weight_grams).font = data_font
                ws.cell(row=row_idx, column=4, value=round(shot.velocity, 2)).font = data_font
                ws.cell(row=row_idx, column=5, value=round(shot.energy_joules, 2)).font = data_font
                ws.cell(row=row_idx, column=6, value=round(shot.compensation_factor, 3)).font = data_font

            # Auto-fit column widths
            for col in ws.columns:
                max_len = max(len(str(cell.value or "")) for cell in col)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

            wb.save(file_path)
            return True
        except Exception as e:
            print(f"Excel export error: {e}")
            return False

    @staticmethod
    def export_pdf(file_path: Path, shots: List[Shot], gun_name: str = "", player_name: str = "", stats: Optional[Dict[str, Any]] = None) -> bool:
        try:
            doc = SimpleDocTemplate(
                str(file_path),
                pagesize=A4,
                leftMargin=36,
                rightMargin=36,
                topMargin=36,
                bottomMargin=36,
            )

            styles = getSampleStyleSheet()
            title_style = ParagraphStyle(
                "TitleStyle",
                parent=styles["Heading1"],
                fontSize=20,
                leading=24,
                textColor=colors.HexColor("#1B5E20"),
                spaceAfter=6,
            )
            sub_style = ParagraphStyle(
                "SubStyle",
                parent=styles["Normal"],
                fontSize=10,
                leading=14,
                textColor=colors.HexColor("#444444"),
            )

            elements = []

            # Header title
            elements.append(Paragraph("<b>ChronoMate Ballistic Report</b>", title_style))
            date_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            meta_info = f"<b>Date:</b> {date_str} &nbsp;|&nbsp; <b>Total Shots:</b> {len(shots)}"
            if gun_name:
                meta_info += f" &nbsp;|&nbsp; <b>Gun:</b> {gun_name}"
            if player_name:
                meta_info += f" &nbsp;|&nbsp; <b>Player:</b> {player_name}"
            elements.append(Paragraph(meta_info, sub_style))
            elements.append(Spacer(1, 14))

            # Stats Table
            if stats:
                stats_data = [
                    ["AVERAGE", "MAX", "MIN", "EXTREME SPREAD", "STD DEV", "ROF"],
                    [
                        f"{stats.get('avg', 0.0):.1f} m/s",
                        f"{stats.get('max', 0.0):.1f} m/s",
                        f"{stats.get('min', 0.0):.1f} m/s",
                        f"{stats.get('es', 0.0):.1f} m/s",
                        f"{stats.get('sd', 0.0):.2f} m/s",
                        f"{stats.get('rof', '0.0')} r/m",
                    ],
                ]
                stats_table = Table(stats_data, colWidths=[85] * 6)
                stats_table.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2E7D32")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, 0), 8),
                    ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                    ("BACKGROUND", (0, 1), (-1, 1), colors.HexColor("#F1F8E9")),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#A5D6A7")),
                    ("TOPPADDING", (0, 0), (-1, -1), 6),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ]))
                elements.append(stats_table)
                elements.append(Spacer(1, 14))

            # Velocity Trend Drawing
            if len(shots) > 1:
                drawing_width = 520
                drawing_height = 110
                dwg = Drawing(drawing_width, drawing_height)

                # Background
                dwg.add(Rect(0, 0, drawing_width, drawing_height, fillColor=colors.HexColor("#F5F5F5"), strokeColor=colors.HexColor("#E0E0E0")))

                vels = [s.velocity for s in shots]
                min_v = min(vels)
                max_v = max(vels)
                range_v = max(1.0, max_v - min_v)
                disp_min = max(0.0, min_v - range_v * 0.2)
                disp_max = max_v + range_v * 0.2
                disp_range = max(0.1, disp_max - disp_min)

                avg_v = sum(vels) / len(vels)
                avg_y = ((avg_v - disp_min) / disp_range) * (drawing_height - 20) + 10
                dwg.add(Line(0, avg_y, drawing_width, avg_y, strokeColor=colors.HexColor("#E65100"), strokeWidth=1, strokeDashArray=[4, 4]))

                # Trend lines and points
                for i in range(len(shots) - 1):
                    x1 = (i / (len(shots) - 1)) * (drawing_width - 20) + 10
                    y1 = ((shots[i].velocity - disp_min) / disp_range) * (drawing_height - 20) + 10
                    x2 = ((i + 1) / (len(shots) - 1)) * (drawing_width - 20) + 10
                    y2 = ((shots[i + 1].velocity - disp_min) / disp_range) * (drawing_height - 20) + 10
                    dwg.add(Line(x1, y1, x2, y2, strokeColor=colors.HexColor("#2E7D32"), strokeWidth=2))
                    dwg.add(Circle(x1, y1, 2.5, fillColor=colors.HexColor("#1B5E20"), strokeColor=colors.transparent))

                last_x = drawing_width - 10
                last_y = ((shots[-1].velocity - disp_min) / disp_range) * (drawing_height - 20) + 10
                dwg.add(Circle(last_x, last_y, 2.5, fillColor=colors.HexColor("#1B5E20"), strokeColor=colors.transparent))

                elements.append(dwg)
                elements.append(Spacer(1, 14))

            # Full Shot Table
            table_data = [["#", "Weight (g)", "Velocity (m/s)", "Energy (J)"]]
            for idx, shot in enumerate(shots, 1):
                table_data.append([
                    f"{idx:02d}",
                    f"{shot.weight_grams:.2f}",
                    f"{shot.velocity:.1f}",
                    f"{shot.energy_joules:.2f}",
                ])

            shot_table = Table(table_data, colWidths=[60, 150, 150, 160])
            shot_table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#424242")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, 0), 9),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F9F9F9")]),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E0E0E0")),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]))

            elements.append(shot_table)
            doc.build(elements)
            return True
        except Exception as e:
            print(f"PDF export error: {e}")
            return False
